"""Calls the real FastAPI app through its HTTP interface (starlette TestClient)."""
import importlib
import io
import os

import pytest
from fastapi.testclient import TestClient

import main as api_main
from conftest import make_stub_detector

REQUIRED = {"file", "modality", "verdict", "confidence_real_pct", "confidence_fake_pct",
            "model_used", "evidence", "warnings", "decision_margin_pct", "confidence_level"}


@pytest.fixture
def client():
    api_main._detectors.clear()
    return TestClient(api_main.app)


def _post(client, path, files, data=None):
    handles = {k: (os.path.basename(v), open(v, "rb")) for k, v in files.items()}
    try:
        return client.post(path, files=handles, data=data or {})
    finally:
        for _, h in handles.values():
            h.close()


# ---------------- _annotate (confidence labels) ----------------
@pytest.mark.parametrize("fake,real,verdict,margin,level", [
    (46.96, 53.04, "REAL", 6.08, "low"),         # the user's 53/47 case: visibly tentative
    (35.0, 65.0, "REAL", 30.0, "medium"),
    (5.0, 95.0, "REAL", 90.0, "high"),
    (90.0, 10.0, "FAKE", 80.0, "high"),
    (50.0, 50.0, "UNCERTAIN", 0.0, "low"),
    (15.88, 84.12, "UNCERTAIN", 68.24, "low"),   # UNCERTAIN is always low, even with a big gap (signature fallback)
    (42.5, 57.5, "REAL", 15.0, "medium"),        # boundary: margin 15 -> medium
    (30.0, 70.0, "REAL", 40.0, "high"),          # boundary: margin 40 -> high
])
def test_annotate_levels(fake, real, verdict, margin, level):
    d = api_main._annotate({"verdict": verdict, "confidence_fake_pct": fake, "confidence_real_pct": real})
    assert d["decision_margin_pct"] == pytest.approx(margin, abs=0.01) and d["confidence_level"] == level

def test_annotate_survives_garbage():
    d = api_main._annotate({"verdict": "REAL", "confidence_fake_pct": None, "confidence_real_pct": "x"})
    assert d["confidence_level"] == "low"


# ---------------- all five endpoints, real files ----------------
def test_health(client):
    assert client.get("/api/health").json()["status"] == "ok"

def test_image_endpoint_users_exact_numbers_end_to_end(client, natural_image):
    api_main._detectors["image"] = make_stub_detector(fake_pct=46.96)
    r = _post(client, "/api/image", {"file": natural_image})
    j = r.json()
    assert r.status_code == 200 and REQUIRED <= set(j)
    assert j["verdict"] == "REAL" and j["confidence_real_pct"] == 53.04
    assert j["confidence_level"] == "low" and j["decision_margin_pct"] == 6.08

def test_image_endpoint_confident_result_is_high(client, natural_image):
    api_main._detectors["image"] = make_stub_detector(fake_pct=2.0)
    j = _post(client, "/api/image", {"file": natural_image}).json()
    assert j["verdict"] == "REAL" and j["confidence_level"] == "high"

def test_image_endpoint_forensic_fallback(client, natural_image):
    j = _post(client, "/api/image", {"file": natural_image}).json()
    assert REQUIRED <= set(j) and "forensic-heuristics-only" in j["model_used"]

def test_document_endpoint(client, pdf_file):
    r = _post(client, "/api/document", {"file": pdf_file})
    j = r.json()
    assert r.status_code == 200 and REQUIRED <= set(j) and j["modality"] == "document"
    assert len(j["sub_results"]) == 1

def test_document_endpoint_leaves_no_temp_files_in_cwd(client, pdf_file, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _post(client, "/api/document", {"file": pdf_file})
    assert [f for f in os.listdir(tmp_path) if f.startswith("__tmp")] == []

def test_audio_endpoint_and_speaker_comparison(client, wav_file):
    r = _post(client, "/api/audio", {"file": wav_file, "reference": wav_file})
    j = r.json()
    assert r.status_code == 200 and REQUIRED <= set(j)
    assert j["speaker_comparison"]["verdict"] == "LIKELY SAME SPEAKER"

def test_video_endpoint(client, video_file):
    r = _post(client, "/api/video", {"file": video_file}, data={"max_frames": "4"})
    j = r.json()
    assert r.status_code == 200 and REQUIRED <= set(j) and j["modality"] == "video"
    assert len(j["sub_results"]) >= 4

def test_signature_endpoint_with_and_without_reference(client, signature_files):
    ref, q = signature_files
    j = _post(client, "/api/signature", {"questioned": q, "reference": ref}).json()
    assert REQUIRED <= set(j) and j["verdict"] in ("GENUINE", "FORGED", "UNCERTAIN")
    j2 = _post(client, "/api/signature", {"questioned": q}).json()
    assert j2["verdict"] == "UNCERTAIN" and j2["confidence_level"] == "low"

def test_garbage_upload_returns_json_error_the_frontend_can_show(client, tmp_path):
    bad = tmp_path / "bad.png"; bad.write_bytes(b"this is not an image")
    r = _post(client, "/api/image", {"file": str(bad)})
    assert r.status_code == 500 and r.headers["content-type"].startswith("application/json")
    assert "message" in r.json()          # api.js reads data.message


# ---------------- CORS ----------------
def _reload_main(monkeypatch, **env):
    for k in ("CORS_ALLOW_ALL", "CORS_ALLOW_ORIGINS"):
        monkeypatch.delenv(k, raising=False)
    for k, v in env.items():
        monkeypatch.setenv(k, v)
    return TestClient(importlib.reload(api_main).app)

def _preflight(c, origin):
    return c.options("/api/image", headers={
        "Origin": origin, "Access-Control-Request-Method": "POST",
        "Access-Control-Request-Headers": "ngrok-skip-browser-warning,content-type"})

def test_cors_default_allows_vite_dev_server_and_the_ngrok_header(monkeypatch):
    c = _reload_main(monkeypatch)
    r = _preflight(c, "http://localhost:5173")
    assert r.status_code == 200 and r.headers["access-control-allow-origin"] == "http://localhost:5173"
    assert "ngrok-skip-browser-warning" in r.headers["access-control-allow-headers"].lower()

def test_cors_default_rejects_unknown_origin(monkeypatch):
    c = _reload_main(monkeypatch)
    r = _preflight(c, "https://evil.example.com")
    assert r.status_code == 400 and "access-control-allow-origin" not in r.headers

def test_cors_extra_origin_via_env(monkeypatch):
    c = _reload_main(monkeypatch, CORS_ALLOW_ORIGINS="https://my-app.example.com, http://192.168.1.5:5173")
    for o in ("https://my-app.example.com", "http://192.168.1.5:5173", "http://localhost:5173"):
        assert _preflight(c, o).headers["access-control-allow-origin"] == o

def test_cors_allow_all_switch(monkeypatch):
    c = _reload_main(monkeypatch, CORS_ALLOW_ALL="1")
    r = _preflight(c, "https://anything.example.org")
    assert r.status_code == 200 and r.headers["access-control-allow-origin"] == "*"
    assert c.get("/api/health", headers={"Origin": "https://x.y"}).headers["access-control-allow-origin"] == "*"
    _reload_main(monkeypatch)       # restore default app state for later tests


# ---------------- other document formats the frontend advertises ----------------
def test_document_endpoint_docx_and_pptx(client, tmp_path):
    import docx
    import pptx
    from PIL import Image
    from skimage import data
    img = tmp_path / "pic.png"
    Image.fromarray(data.coffee()).save(img)

    d = docx.Document(); d.add_paragraph("hello"); d.add_picture(str(img)); dp = tmp_path / "t.docx"; d.save(str(dp))
    pr = pptx.Presentation(); s = pr.slides.add_slide(pr.slide_layouts[6]); s.shapes.add_picture(str(img), 0, 0)
    pp = tmp_path / "t.pptx"; pr.save(str(pp))

    for path in (dp, pp):
        r = _post(client, "/api/document", {"file": str(path)})
        j = r.json()
        assert r.status_code == 200, j
        assert REQUIRED <= set(j) and len(j["sub_results"]) == 1, path


# ---------------- regressions found while reviewing the redesigned frontend ----------------
@pytest.mark.parametrize("n", ["0", "-5", "500"])
def test_video_endpoint_survives_silly_frame_counts(client, video_file, n):
    """Typing 0 into the 'frames to sample' box used to crash the server with HTTP 500."""
    r = _post(client, "/api/video", {"file": video_file}, data={"max_frames": n})
    assert r.status_code == 200, r.text


def test_verdict_words_the_frontend_must_understand(client, natural_image, signature_files):
    """The UI maps verdict words to colours. Signatures use GENUINE/FORGED, everything else REAL/FAKE.
    If these change, 08_frontend/src/lib/types.js (verdictView) must change too."""
    assert _post(client, "/api/image", {"file": natural_image}).json()["verdict"] in {"REAL", "FAKE", "UNCERTAIN"}
    ref, questioned = signature_files[0], signature_files[1]
    r = _post(client, "/api/signature", {"questioned": questioned, "reference": ref})
    assert r.json()["verdict"] in {"GENUINE", "FORGED", "UNCERTAIN"}


# ---------------- "which model is live?" must be visible, and failures must explain themselves ----------------
def test_health_reports_pretrained_model_by_default(client, monkeypatch, tmp_path):
    import image_detector as imd
    monkeypatch.delenv("IMAGE_MODEL_ID", raising=False)
    monkeypatch.setattr(imd, "FINETUNED_DIR", str(tmp_path / "none"))
    j = client.get("/api/health").json()
    assert j["status"] == "ok" and j["image_model_is_finetuned"] is False
    assert j["image_model"] == imd.DEFAULT_MODEL_ID and "pretrained" in j["image_model_note"].lower()
    assert "face_detection" in j and "face_detection_reason" in j


def test_health_reports_finetuned_model_when_installed(client, monkeypatch, tmp_path):
    import image_detector as imd
    d = tmp_path / "image_finetuned"; d.mkdir(); (d / "config.json").write_text("{}")
    monkeypatch.delenv("IMAGE_MODEL_ID", raising=False)
    monkeypatch.setattr(imd, "FINETUNED_DIR", str(d))
    j = client.get("/api/health").json()
    assert j["image_model_is_finetuned"] is True and j["image_model"] == "image_finetuned"


def test_face_detection_failure_names_the_real_reason(monkeypatch):
    import sys, face_crop
    monkeypatch.setattr(face_crop, "_cascade", None)
    monkeypatch.setattr(face_crop, "_cascade_failed", False)
    monkeypatch.setattr(face_crop, "_failure_reason", "")
    monkeypatch.setitem(sys.modules, "cv2", None)               # makes "import cv2" fail
    ok, why = face_crop.face_detection_status()
    assert ok is False and "could not be imported" in why


def test_prediction_warning_includes_the_face_detection_reason(natural_image, monkeypatch):
    import sys, face_crop
    from conftest import make_stub_detector
    monkeypatch.setattr(face_crop, "_cascade", None)
    monkeypatch.setattr(face_crop, "_cascade_failed", False)
    monkeypatch.setattr(face_crop, "_failure_reason", "")
    monkeypatch.setitem(sys.modules, "cv2", None)
    monkeypatch.setenv("DEEPFAKE_CALIBRATION", "off")
    r = make_stub_detector(fake_pct=60.0, face_crop=True, fusion=False).predict(natural_image)
    text = " ".join(r.warnings)
    assert "Face detection unavailable" in text and "could not be imported" in text


def test_uncertain_band_can_be_configured_but_defaults_to_six(monkeypatch):
    from common.schema import verdict_from_scores
    monkeypatch.delenv("UNCERTAIN_BAND", raising=False)
    assert verdict_from_scores(49, 51) == "UNCERTAIN"                 # default band 6
    monkeypatch.setenv("UNCERTAIN_BAND", "0")
    assert verdict_from_scores(49, 51) == "REAL"                      # never answers UNCERTAIN
    monkeypatch.setenv("UNCERTAIN_BAND", "not-a-number")
    assert verdict_from_scores(49, 51) == "UNCERTAIN"                 # bad value falls back safely
